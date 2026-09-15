import re
import unicodedata
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from pydantic import BaseModel, ConfigDict, Field
from pypdf import PdfReader

from app.services.ats_keyword_service import extract_jd_keywords
from app.services.resume_service import ResumeServiceError


router = APIRouter(
    prefix="/api/ats",
    tags=["WAYVORA Resume Analysis"],
)

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_DOCUMENT_BYTES = 25 * 1024 * 1024
MAX_RESUME_CHARACTERS = 50000


class ATSRequest(BaseModel):
    model_config = ConfigDict(
        extra="forbid",
        str_strip_whitespace=True,
    )

    resume_text: str = Field(
        min_length=50,
        max_length=MAX_RESUME_CHARACTERS,
    )

    job_description: str = Field(
        min_length=30,
        max_length=30000,
    )

    keywords: list[str] = Field(
        min_length=1,
        max_length=60,
    )


SECTION_NAMES = {
    "summary": {
        "summary",
        "professional summary",
        "profile",
        "professional profile",
        "career objective",
        "objective",
    },
    "skills": {
        "skills",
        "technical skills",
        "core skills",
        "core competencies",
        "key skills",
        "skills and technologies",
    },
    "experience": {
        "experience",
        "work experience",
        "professional experience",
        "employment history",
        "internships",
        "internship experience",
    },
    "projects": {
        "projects",
        "academic projects",
        "personal projects",
        "selected projects",
    },
    "education": {
        "education",
        "educational qualifications",
        "academic qualifications",
        "academic background",
    },
    "certifications": {
        "certifications",
        "certificates",
        "licenses and certifications",
        "certifications and training",
    },
}


def normalise(text: str) -> str:
    text = unicodedata.normalize("NFKC", text).casefold()
    text = text.replace("\u00ad", "")
    text = re.sub(r"[\u2010-\u2015]", "-", text)
    return re.sub(r"\s+", " ", text).strip()


def contains_phrase(text: str, phrase: str) -> bool:
    # Avoid partial-word matches, such as Java inside JavaScript.
    pattern = r"(?<!\w)" + re.escape(phrase) + r"(?!\w)"
    return re.search(pattern, text) is not None


def section_checks(resume_text: str) -> dict[str, bool]:
    headings = {
        normalise(line).strip(" :.-•*")
        for line in resume_text.splitlines()
        if line.strip()
    }

    return {
        section: bool(headings.intersection(names))
        for section, names in SECTION_NAMES.items()
    }


def extract_resume_document(
    filename: str,
    data: bytes,
) -> tuple[str, list[str]]:
    extension = Path(filename).suffix.lower()
    warnings = []

    try:
        if extension == ".pdf":
            reader = PdfReader(BytesIO(data))

            if reader.is_encrypted:
                raise HTTPException(
                    status_code=422,
                    detail="Upload a PDF without password protection.",
                )

            if len(reader.pages) > 20:
                raise HTTPException(
                    status_code=422,
                    detail="Upload a resume with no more than 20 pages.",
                )

            parts = []
            empty_pages = []

            for index, page in enumerate(reader.pages, start=1):
                page_text = page.extract_text() or ""
                parts.append(page_text)

                if not page_text.strip():
                    empty_pages.append(str(index))

            text = "\n\n".join(parts)

            if empty_pages:
                warnings.append(
                    "No text was extracted from PDF page(s): "
                    + ", ".join(empty_pages)
                    + ". These pages may be blank or scanned."
                )

        elif extension == ".docx":
            with ZipFile(BytesIO(data)) as archive:
                expanded_size = sum(
                    item.file_size for item in archive.infolist()
                )

                if expanded_size > MAX_DOCUMENT_BYTES:
                    raise HTTPException(
                        status_code=422,
                        detail="This DOCX is too large after decompression.",
                    )

            document = Document(BytesIO(data))
            parts = [
                paragraph.text
                for paragraph in document.paragraphs
            ]

            def add_tables(tables):
                for table in tables:
                    for row in table.rows:
                        for cell in row.cells:
                            parts.extend(
                                paragraph.text
                                for paragraph in cell.paragraphs
                            )
                            add_tables(cell.tables)

            add_tables(document.tables)

            for section in document.sections:
                areas = (
                    section.header,
                    section.first_page_header,
                    section.even_page_header,
                    section.footer,
                    section.first_page_footer,
                    section.even_page_footer,
                )

                for area in areas:
                    parts.extend(
                        paragraph.text
                        for paragraph in area.paragraphs
                    )
                    add_tables(area.tables)

            text = "\n".join(parts)

            warnings.append(
                "Check the extracted text. Text boxes, images and some "
                "document elements may not be included; table text may "
                "appear in a different order."
            )

        else:
            raise HTTPException(
                status_code=415,
                detail=(
                    "Upload a PDF or DOCX file. "
                    "Old .doc files are not supported."
                ),
            )

    except HTTPException:
        raise

    except Exception as exc:
        raise HTTPException(
            status_code=422,
            detail=(
                "This document could not be read. Try exporting it again "
                "as a text-based PDF or DOCX."
            ),
        ) from exc

    text = text.replace("\x00", "").strip()

    if len(text) < 50:
        raise HTTPException(
            status_code=422,
            detail=(
                "Not enough readable text was extracted. "
                "For scanned resumes, use OCR or upload a text-based "
                "PDF or DOCX."
            ),
        )

    if len(text) > MAX_RESUME_CHARACTERS:
        raise HTTPException(
            status_code=422,
            detail="The extracted resume exceeds 50,000 characters.",
        )

    warnings.append(
        "Review the extracted text before analysing. Successful extraction "
        "here does not guarantee identical parsing by every employer's ATS."
    )

    return text, warnings


@router.post("/prepare-upload")
def prepare_resume_upload(
    resume_file: UploadFile = File(...),
    job_description: str = Form(
        ...,
        min_length=30,
        max_length=30000,
    ),
):
    try:
        job_description = job_description.strip()

        if len(job_description) < 30:
            raise HTTPException(
                status_code=422,
                detail=(
                    "Paste a job description containing at least "
                    "30 characters."
                ),
            )

        data = resume_file.file.read(MAX_UPLOAD_BYTES + 1)

        if not data:
            raise HTTPException(
                status_code=422,
                detail="The uploaded file is empty.",
            )

        if len(data) > MAX_UPLOAD_BYTES:
            raise HTTPException(
                status_code=413,
                detail="Upload a file smaller than 5 MB.",
            )

        text, warnings = extract_resume_document(
            resume_file.filename or "",
            data,
        )

        try:
            keywords, keyword_warnings = extract_jd_keywords(
                job_description
            )

        except ResumeServiceError as exc:
            if exc.status_code == 413:
                message = (
                    "The job description is too long for Qwen's current "
                    "context limit. Keep the responsibilities and "
                    "requirements, then try again. No text was "
                    "shortened automatically."
                )

            elif exc.status_code == 503:
                message = (
                    "Qwen is unavailable or busy. Make sure your "
                    "local model is running on port 8081, then try again."
                )

            elif exc.status_code == 504:
                message = (
                    "Qwen took too long to extract keywords. "
                    "Try a shorter job description."
                )

            else:
                message = exc.message

            raise HTTPException(
                status_code=exc.status_code,
                detail=message,
            ) from exc

        warnings.extend(keyword_warnings)

        return {
            "resume_text": text,
            "suggested_keywords": keywords,
            "warnings": warnings,
        }

    finally:
        resume_file.file.close()


@router.post("/analyze")
def analyze_resume(payload: ATSRequest):
    keywords = []
    seen = set()

    for value in payload.keywords:
        label = value.strip()
        key = normalise(label)

        if not key:
            continue

        if len(key) > 100:
            raise HTTPException(
                status_code=422,
                detail="Each keyword must be 100 characters or fewer.",
            )

        if key not in seen:
            keywords.append((label, key))
            seen.add(key)

    if not keywords:
        raise HTTPException(
            status_code=422,
            detail="Enter at least one keyword from the job description.",
        )

    resume = normalise(payload.resume_text)
    job_description = normalise(payload.job_description)

    not_in_jd = [
        label
        for label, key in keywords
        if not contains_phrase(job_description, key)
    ]

    if not_in_jd:
        raise HTTPException(
            status_code=422,
            detail=(
                "These keywords were not found in the job description: "
                + ", ".join(not_in_jd)
                + ". Copy the wording from the job description."
            ),
        )

    matched = []
    missing = []

    for label, key in keywords:
        if contains_phrase(resume, key):
            matched.append(label)
        else:
            missing.append(label)

    score = round(
        100 * len(matched) / len(keywords),
        1,
    )

    sections = section_checks(payload.resume_text)
    suggestions = []

    if missing:
        suggestions.append(
            "Review the keywords not found. Add a term only if your "
            "actual experience supports it, preferably in a relevant bullet."
        )

    if not sections["skills"]:
        suggestions.append(
            "A separate Skills heading was not detected. "
            "Check that your skills are clearly labelled."
        )

    if not sections["education"]:
        suggestions.append(
            "A separate Education heading was not detected. "
            "Check whether an education section is appropriate for this role."
        )

    if not sections["experience"] and not sections["projects"]:
        suggestions.append(
            "Neither an Experience nor a Projects heading was detected. "
            "Include relevant work, internships or projects where applicable."
        )

    suggestions.append(
        "Support important skills with truthful examples of how you used "
        "them. Repeating a keyword does not increase this score."
    )

    return {
        "score": score,
        "score_name": "Reviewed-keyword coverage",
        "matched_count": len(matched),
        "keyword_count": len(keywords),
        "matched": matched,
        "missing": missing,
        "sections": sections,
        "suggestions": suggestions,
        "method": (
            "Matched unique keywords divided by reviewed unique keywords, "
            "multiplied by 100. All reviewed keywords have equal weight."
        ),
        "limitations": [
            "This is a WAYVORA estimate, not an employer's ATS score.",
            "Qwen can miss or misidentify requirements. "
            "Only the reviewed keywords are measured.",
            "Matching ignores case and repeated whitespace, "
            "but does not infer synonyms or proficiency.",
            "A keyword mention is not proof of experience or eligibility.",
            "Successful extraction does not guarantee that another "
            "system will read the document identically.",
            "Section checks detect standalone headings, "
            "not content quality or visual formatting.",
        ],
    }